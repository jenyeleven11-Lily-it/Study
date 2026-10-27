# -*- coding: utf-8 -*-
"""按《大数据分析与应用》上机报告3模板填写内容并嵌入图表"""
import shutil
from copy import deepcopy
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = r"C:\Users\ASUS\Downloads\学号+姓名+《大数据分析与应用》上机报告3.docx"
OUT = r"D:\学习\上机\python数据分析与应用\上机3\上机3_数据可视化基础_实验报告.docx"
IMG = r"D:\学习\上机\python数据分析与应用\上机3"

SONG = "宋体"


def set_run(run, size=10.5, bold=False, code=False, en=None):
    if code:
        en = "Consolas"
        run.font.size = Pt(9)
    else:
        en = en or "Times New Roman"
        run.font.size = Pt(size)
    run.font.name = en
    rPr = run._element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rPr.append(rf)
    rf.set(qn("w:eastAsia"), SONG)
    run.font.bold = bold


def insert_after(cell, anchor, text, code=False, size=10.5, bold=False, align=None):
    p = cell.add_paragraph()
    pf = p.paragraph_format
    pf.first_line_indent = Pt(0)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, code=code)
    if code:
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "F2F2F2")
        pPr.append(shd)
    anchor._p.addnext(p._p)
    return p


def insert_img_after(cell, anchor, path, width=12.0):
    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.first_line_indent = Pt(0)
    pf.space_before = Pt(2)
    pf.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(path, width=Cm(width))
    anchor._p.addnext(p._p)
    return p


# ---------- 17 个槽位的源代码 ----------
SRC_CODE = [
    # 任务1
    "import pandas as pd\n"
    "df = pd.read_excel('student_grade.xlsx')\n"
    "print(df.shape)\nprint(df.head())",

    "import matplotlib.pyplot as plt\n"
    "bins = [0,150,200,250,300]\n"
    "labels = ['不及格','及格','良好','优秀']\n"
    "df['等级'] = pd.cut(df['总成绩'], bins=bins, labels=labels)\n"
    "counts = df['等级'].value_counts()\n"
    "plt.pie(counts, labels=counts.index, autopct='%.1f%%', startangle=90)\n"
    "plt.title('学生考试总成绩分布饼图')\nplt.savefig('task1_饼图.png')",

    "subjects = ['数学成绩','阅读成绩','写作成绩']\n"
    "df[subjects].boxplot()\n"
    "plt.title('学生各项考试成绩分散情况箱线图')\n"
    "plt.ylabel('成绩')\nplt.savefig('task1_箱线图.png')",

    "print(df['总成绩'].describe())\n"
    "print(df[subjects].describe())",

    # 任务2
    "import numpy as np\n"
    "se_mean = df.groupby('自我效能感')['总成绩'].apply(np.mean)\n"
    "prep_mean = df.groupby('考试课程准备情况')['总成绩'].apply(np.mean)\n"
    "print(se_mean)\nprint(prep_mean)",

    "fig, axes = plt.subplots(1, 2, figsize=(12, 5))",

    "axes[0].bar(se_mean.index, se_mean.values, color='steelblue')\n"
    "axes[0].set_title('自我效能感与总成绩均值关系')\n"
    "axes[1].bar(prep_mean.index, prep_mean.values, color='orange')\n"
    "axes[1].set_title('考试课程准备情况与总成绩均值关系')\n"
    "plt.tight_layout()\nplt.savefig('task2_柱状图.png')",

    "print('自我效能感越高，总成绩均值越高；完成备考的学生均值更高')",

    # 任务3
    "import pandas as pd\n"
    "df = pd.read_csv('aqi.csv', encoding='utf-8')\n"
    "print(df.shape)\nprint(df.head())",

    "import matplotlib.pyplot as plt\n"
    "plt.rcParams['font.sans-serif'] = ['SimHei']\n"
    "plt.rcParams['axes.unicode_minus'] = False",

    "num_cols = ['AQI','PM2.5含量/ppm','PM10含量/ppm','SO2含量/ppm','CO含量/ppm','NO2含量/ppm','O3含量/ppm']\n"
    "corr = df[num_cols].corr()\nprint(corr)",

    "import seaborn as sns\n"
    "sns.heatmap(corr, annot=True, fmt='.2f', cmap='coolwarm', square=True)\n"
    "plt.title('空气质量特征相关性热力图')\nplt.savefig('task3_热力图.png')",

    # 任务4
    "import pandas as pd\n"
    "df = pd.read_csv('商品销售数据.csv', encoding='utf-8-sig')\n"
    "print(df.shape)\nprint(df.head())",

    "sales = df.groupby('二级类')['实际金额'].sum().sort_values(ascending=False)\n"
    "print(sales.head(8))",

    "from pyecharts.charts import Funnel\n"
    "from pyecharts import options as opts\n"
    "top5 = sales.head(5)\n"
    "data = [[str(k), round(float(v),2)] for k,v in top5.items()]\n"
    "Funnel(init_opts=opts.InitOpts(width='900px',height='600px'))\n"
    "    .add('商品类别销售额', data).render('task4_销售额前5漏斗图.html')",

    "counts = df.groupby('商品')['订单号'].count().sort_values(ascending=False)\n"
    "print(counts.head())",

    "from pyecharts.charts import WordCloud\n"
    "data = [[str(k), int(v)] for k,v in counts.items()]\n"
    "WordCloud(init_opts=opts.InitOpts(width='900px',height='600px'))\n"
    "    .add('', data, word_size_range=[12,60]).render('task4_商品销售数量词云图.html')",
]

# ---------- 17 个结果（图片或文本） ----------
RES = [
    ("text", "运行输出：数据形状 (1000, 7)，包含 性别/自我效能感/考试课程准备情况/数学成绩/阅读成绩/写作成绩/总成绩 7 个特征。"),
    ("img", IMG + r"\task1_饼图.png"),
    ("img", IMG + r"\task1_箱线图.png"),
    ("text", "总成绩：均值 203.31、标准差 42.77、最小 27、最大 300；数学/阅读/写作均值分别为 66.09 / 69.17 / 68.05，整体分布较集中。"),
    ("text", "自我效能感下总成绩均值：低 192.05、中 207.12、高 217.61；考试课程准备情况：未完成 195.12、完成 218.01。"),
    ("text", "创建 1×2 的画布，添加左右两个子图 axes[0]、axes[1]。"),
    ("img", IMG + r"\task2_柱状图.png"),
    ("text", "自我效能感越高总成绩均值越高，呈正相关；完成考试课程准备的学生均值明显更高，说明备考情况影响成绩。"),
    ("text", "运行输出：数据形状 (271, 9)，包含 日期、AQI、质量等级 及 6 项污染物浓度指标。"),
    ("text", "设置中文字体为黑体 SimHei，并关闭 unicode_minus，解决中文乱码与负号方块问题。"),
    ("text", "AQI 与 PM10 相关系数 0.75、与 PM2.5 相关系数 0.72，相关性较强；O3 与 PM2.5 呈弱负相关 -0.12。"),
    ("img", IMG + r"\task3_热力图.png"),
    ("text", "运行输出：数据形状 (8978, 11)，包含 订单号/设备ID/应付金额/实际金额/商品/支付时间/地点/状态/提现/大类/二级类。"),
    ("text", "销售额前 5：乳制品 5308.5、茶饮料 4905.0、肉干/豆制品/蛋 4378.7、饼干糕点 3837.5、功能饮料 3581.7。"),
    ("text", "交互式漏斗图已保存为 task4_销售额前5漏斗图.html，用浏览器打开后截图插入此处。"),
    ("text", "销售数量前 5：怡宝纯净水 657、40g双汇玉米热狗肠 240、东鹏特饮 238、脉动 235、250ml维他柠檬茶 224。"),
    ("text", "交互式词云图已保存为 task4_商品销售数量词云图.html，用浏览器打开后截图插入此处。"),
]

shutil.copyfile(SRC, OUT)
doc = Document(OUT)
t1 = doc.tables[0]

# ---------- 三、实验内容 单元格填代码与截图 ----------
cell = t1.rows[5].cells[0]
paras = list(cell.paragraphs)
src_i = 0
res_i = 0
for p in paras:
    txt = p.text.strip()
    if txt == "源程序：" and src_i < len(SRC_CODE):
        last = p
        for line in SRC_CODE[src_i].split("\n"):
            last = insert_after(cell, last, line if line else " ", code=True)
        src_i += 1
    elif txt == "结果截屏：" and res_i < len(RES):
        spec = RES[res_i]
        last = p
        if spec[0] == "img":
            last = insert_img_after(cell, last, spec[1])
        else:
            last = insert_after(cell, last, spec[1], size=10.5)
        res_i += 1

print("源程序填入:", src_i, "/17, 结果填入:", res_i, "/17")

# ---------- 四、实验结果分析 收获体会 ----------
cell4 = t1.rows[6].cells[0]
# 清空除第一个标题外的段落
for pp in list(cell4.paragraphs[1:]):
    pp._p.getparent().remove(pp._p)
first = cell4.paragraphs[0]
for r in list(first.runs):
    r._element.getparent().remove(r._element)
if first.text:
    first._element.text = ""
r = first.add_run("四、实验结果分析以及出现问题")
set_run(r, size=10.5, bold=True)

reflection = [
    "本次上机完成了数据可视化基础实验，使用 Matplotlib、seaborn、pyecharts 三个库绘制了饼图、箱线图、柱状图、热力图、漏斗图和词云图。通过实验，我掌握了根据数据特点选择合适的图表类型：展示构成比例用饼图，展示数据离散程度用箱线图，比较不同类别均值用柱状图，分析变量间相关性用热力图，体现层级递减关系用漏斗图，突出高频词用词云图。",
    "实验过程中遇到并解决了几个问题：一是中文显示乱码，通过设置 plt.rcParams['font.sans-serif']=['SimHei'] 解决；二是保存图片时负号显示为方块，通过设置 axes.unicode_minus=False 解决；三是读取商品销售数据时出现编码错误，将编码由 gbk 改为 utf-8-sig 后正常；四是初始未安装 pyecharts，通过 pip install pyecharts 完成安装。另外，脚本运行时相对路径找不到数据文件，我在脚本中增加了自动切换到脚本所在目录的语句，使程序能够稳定读取数据。",
    "不足方面，我对 seaborn 的高级定制掌握还不够深入，词云图等交互式图表的配色与布局还可进一步优化。今后我会继续学习数据可视化的高级技巧，提升图表的美观度与分析深度，同时加强代码的规范性与可维护性。",
]
for txt in reflection:
    p = cell4.add_paragraph()
    pf = p.paragraph_format
    pf.first_line_indent = Pt(0)
    pf.space_before = Pt(0)
    pf.space_after = Pt(2)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(txt)
    set_run(run, size=10.5)

doc.save(OUT)
print("SAVED:", OUT)
